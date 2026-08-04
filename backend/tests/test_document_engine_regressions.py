from __future__ import annotations

import io
import struct
import sys
import unittest
import zlib
from types import SimpleNamespace
from unittest.mock import ANY, patch

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook

from app.models.task import OutputMode
from app.services.document_engine import (
    TranslationContext,
    _replace_docx_images,
    translate_pdf_inplace,
    translate_xls,
    translate_xlsx,
)


class _PrefixTranslator:
    def translate(self, text, target_lang, source_lang, **kwargs):
        return f"译:{text}"


def _png_bytes(color: str) -> bytes:
    rgb = {"red": (255, 0, 0), "blue": (0, 0, 255)}[color]

    def chunk(kind: bytes, payload: bytes) -> bytes:
        body = kind + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

    width = height = 32
    scanlines = b"".join(b"\0" + bytes(rgb) * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(scanlines))
        + chunk(b"IEND", b"")
    )


class DocumentEngineRegressionTests(unittest.TestCase):
    def test_xlsx_formula_tokenizer_does_not_treat_quoted_a1_as_reference(self):
        wb = Workbook()
        ws = wb.active
        ws["B2"] = "B2"
        ws["C1"] = "B2"
        ws["D1"] = "律师"
        ws["D2"] = "普通员工"
        ws["E1"] = '=IF(C1="B2","yes","no")'
        ws["E2"] = '=COUNTIF(D1:D2,"律师")'
        source = io.BytesIO()
        wb.save(source)

        result = translate_xlsx(
            source.getvalue(),
            _PrefixTranslator(),
            TranslationContext(target_lang="zh", output_mode=OutputMode.PLAIN),
        )
        translated = load_workbook(io.BytesIO(result), data_only=False).active

        # "B2" in the formula is a string token, not a reference to cell B2.
        self.assertEqual(translated["B2"].value, "译:B2")
        # A literal used against a real referenced range must still remain stable.
        self.assertEqual(translated["C1"].value, "B2")
        self.assertEqual(translated["D1"].value, "律师")
        self.assertEqual(translated["D2"].value, "译:普通员工")
        self.assertEqual(translated["E1"].value, '=IF(C1="B2","yes","no")')

    def test_docx_ocr_visits_header_footer_images_and_deduplicates_parts(self):
        shared_image = _png_bytes("red")
        footer_image = _png_bytes("blue")
        doc = Document()
        doc.add_picture(io.BytesIO(shared_image))
        doc.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(shared_image))
        doc.sections[0].footer.paragraphs[0].add_run().add_picture(io.BytesIO(footer_image))

        seen: list[bytes] = []

        def fake_ocr(data, *args, **kwargs):
            seen.append(data)
            return data

        fake_ocr_module = SimpleNamespace(process_image_ocr=fake_ocr)
        with patch.dict(sys.modules, {"app.services.ocr": fake_ocr_module}):
            _replace_docx_images(doc, TranslationContext(target_lang="zh", translate_images="yes"))

        # The shared body/header image is one part; the footer has a second part.
        self.assertEqual(len(seen), 2)

    def test_pdf_redaction_preserves_overlapping_image_and_vector_without_pixmap_sampling(self):
        source_doc = fitz.open()
        page = source_doc.new_page()
        page.draw_rect(fitz.Rect(55, 45, 310, 90), color=(1, 0, 0), fill=(0.9, 0.9, 0.9))
        page.insert_image(fitz.Rect(65, 55, 95, 85), stream=_png_bytes("blue"))
        original = "This is a sufficiently long legal sentence for translation."
        page.insert_text((70, 72), original, fontsize=11)
        image_count = len(page.get_images(full=True))
        drawing_count = len(page.get_drawings())
        source = source_doc.tobytes()
        source_doc.close()

        ctx = TranslationContext(target_lang="en", translate_images="no")
        with patch.object(fitz.Page, "get_pixmap", side_effect=AssertionError("unexpected rasterization")):
            result = translate_pdf_inplace(source, _PrefixTranslator(), ctx)

        translated_doc = fitz.open(stream=result, filetype="pdf")
        translated_page = translated_doc[0]
        self.assertNotIn(original, translated_page.get_text())
        self.assertEqual(len(translated_page.get_images(full=True)), image_count)
        self.assertGreaterEqual(len(translated_page.get_drawings()), drawing_count)
        translated_doc.close()

    def test_xls_uses_limited_libreoffice_conversion_before_xlsx_translation(self):
        converted = b"converted-xlsx"
        translated = b"translated-xlsx"
        ctx = TranslationContext(target_lang="zh")
        with (
            patch("app.services.document_engine._libreoffice_convert", return_value=converted) as convert,
            patch("app.services.document_engine.translate_xlsx", return_value=translated) as xlsx,
        ):
            result = translate_xls(b"legacy-xls", _PrefixTranslator(), ctx)

        self.assertEqual(result, translated)
        convert.assert_called_once_with(b"legacy-xls", "xls", "xlsx")
        xlsx.assert_called_once_with(converted, ANY, ctx)


if __name__ == "__main__":
    unittest.main()
