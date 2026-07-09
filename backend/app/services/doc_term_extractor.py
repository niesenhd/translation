"""文档级术语抽取（功能B）：翻译前从全文抽出需统一译法的关键术语。

解决法律文书"同一术语全文译法不一致"——抽取结果合并进该文档 glossary（STRICT），
逐段注入，保证定义术语/专有名词全文统一。一次额外模型调用，廉价。
"""
from __future__ import annotations

import logging

from app.services.translator import Translator

logger = logging.getLogger(__name__)

# 触发抽取的最小全文长度：过短的文档不值得多一次调用
_MIN_TEXT_LEN = 1500


def extract_plain_text(ext: str, data: bytes) -> str:
    """从各格式文件抽取纯文本（用于术语抽取的输入），best-effort。

    复用 document_engine 既有的段落收集器/解码逻辑，保证与实际翻译口径一致。
    不支持的格式返回空串（跳过抽取，不影响翻译）。
    """
    ext = ext.lower().lstrip(".")
    try:
        if ext in ("txt", "md"):
            from app.services.document_engine import _decode_text_bytes
            return _decode_text_bytes(data)

        if ext in ("docx", "doc"):
            import io
            from docx import Document
            from app.services.document_engine import (
                _collect_docx_paragraphs, _extract_docx_paragraph_text,
                _libreoffice_convert,
            )
            docx_data = _libreoffice_convert(data, "doc", "docx") if ext == "doc" else data
            doc = Document(io.BytesIO(docx_data))
            paras = _collect_docx_paragraphs(doc)
            return "\n".join(_extract_docx_paragraph_text(p) for p in paras)

        if ext == "pdf":
            import fitz
            out = []
            with fitz.open(stream=data, filetype="pdf") as doc:
                for page in doc:
                    out.append(page.get_text())
            return "\n".join(out)

        if ext in ("xlsx", "xls", "csv"):
            # 表格类：单元格文本拼接，量通常不大
            import io
            if ext == "csv":
                from app.services.document_engine import _decode_text_bytes
                return _decode_text_bytes(data)
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            cells = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str) and v.strip():
                            cells.append(v)
            return "\n".join(cells)

        if ext in ("pptx", "ppt"):
            import io
            from pptx import Presentation
            from app.services.document_engine import _libreoffice_convert
            pptx_data = _libreoffice_convert(data, "ppt", "pptx") if ext == "ppt" else data
            prs = Presentation(io.BytesIO(pptx_data))
            texts = []

            def _walk(shape):
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            texts.append(cell.text)

            for slide in prs.slides:
                for shape in slide.shapes:
                    _walk(shape)
            return "\n".join(texts)
    except Exception as exc:  # noqa: BLE001
        logger.warning("抽取全文失败（%s），跳过文档术语抽取：%s", ext, exc)
        return ""

    return ""


def extract_document_terms(ext: str, data: bytes, source_lang: str, target_lang: str,
                           translator: Translator) -> list[dict]:
    """抽取本文档的关键术语表 [{source_term, target_term, priority:'strict'}]。

    文本过短或抽取失败时返回空列表（调用方据此决定是否合并）。
    """
    text = extract_plain_text(ext, data)
    if len(text.strip()) < _MIN_TEXT_LEN:
        return []
    if not hasattr(translator, "extract_document_terms"):
        return []
    terms = translator.extract_document_terms(text, target_lang)  # type: ignore[attr-defined]
    if terms:
        logger.info("文档术语抽取：从 %d 字符全文抽出 %d 个统一术语", len(text), len(terms))
    return terms
