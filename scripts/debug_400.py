"""找出导致 400 错误的段落内容"""
import sys
sys.path.insert(0, '/app')

from io import BytesIO
from minio import Minio
from app.core.database import SessionLocal
from app.models.task import TranslationTask
from app.core.config import get_settings
from docx import Document

settings = get_settings()
db = SessionLocal()
task = db.get(TranslationTask, 'ae9cc53f-f275-46b3-88b9-5c58aac29288')

mc = Minio(settings.minio_endpoint, access_key=settings.minio_access_key,
           secret_key=settings.minio_secret_key, secure=False)

obj = mc.get_object(settings.minio_bucket, 'sources/ae9cc53f-f275-46b3-88b9-5c58aac29288.docx')
data = obj.read()
print(f'File size: {len(data)} bytes')

doc = Document(BytesIO(data))

# 收集所有段落（含表格）
all_texts = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        all_texts.append(t)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                t = p.text.strip()
                if t:
                    all_texts.append(t)

print(f'Total non-empty paragraphs: {len(all_texts)}')

# 分析长度分布
lengths = sorted([len(t) for t in all_texts], reverse=True)
print(f'Top 10 longest: {lengths[:10]}')
print(f'> 10000 chars: {sum(1 for l in lengths if l > 10000)}')
print(f'> 50000 chars: {sum(1 for l in lengths if l > 50000)}')
print(f'> 100000 chars: {sum(1 for l in lengths if l > 100000)}')

# 找出可疑段落：有字母但可能触发 400 的
print('\n--- Suspicious paragraphs ---')
count = 0
for i, t in enumerate(all_texts):
    has_alpha = any(c.isalpha() for c in t)
    if not has_alpha:
        continue
    # 可能的问题：太长、或特殊编码
    if len(t) > 30000:
        count += 1
        print(f'  [{i}] len={len(t)} chars: {repr(t[:80])}')
        if count >= 10:
            break

# 检查是否有空字节或控制字符
print('\n--- Paragraphs with control chars ---')
ctrl_count = 0
for i, t in enumerate(all_texts):
    if any(ord(c) < 32 and c not in '\n\r\t' for c in t):
        ctrl_count += 1
        if ctrl_count <= 5:
            print(f'  [{i}] len={len(t)}: {repr(t[:80])}')
print(f'Total with control chars: {ctrl_count}')

# 实际测试 API 调用
print('\n--- Testing API on first 5 paragraphs ---')
import openai
client = openai.OpenAI(api_key=settings.dashscope_api_key, base_url=settings.dashscope_base_url, max_retries=0)
for i in range(min(5, len(all_texts))):
    t = all_texts[i]
    try:
        r = client.chat.completions.create(
            model='qwen3.6-27b',
            messages=[{'role':'user','content':t}],
            temperature=0.2,
            extra_body={'enable_thinking': False},
        )
        print(f'  [{i}] OK len_in={len(t)} len_out={len(r.choices[0].message.content or "")}')
    except Exception as e:
        print(f'  [{i}] FAIL len_in={len(t)}: {str(e)[:200]}')

db.close()
