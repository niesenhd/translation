"""Debug .doc paragraph extraction"""
import minio
import subprocess
import sys
import os

sys.path.insert(0, "/app")
from app.core.config import get_settings

s = get_settings()
mc = minio.Minio(s.minio_endpoint, s.minio_access_key, s.minio_secret_key, secure=False)

TASK_ID = "117b4155-769d-4e89-9972-7adb9de6de82"

# Download source .doc
obj = mc.get_object(s.minio_bucket, f"sources/{TASK_ID}.doc")
data = obj.read()
with open("/tmp/source.doc", "wb") as f:
    f.write(data)
print(f"Source .doc size: {len(data)} bytes")

# Convert to docx
subprocess.run(
    ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", "/tmp", "/tmp/source.doc"],
    capture_output=True, timeout=60
)

# Parse
from docx import Document
doc = Document("/tmp/source.docx")
paras = doc.paragraphs
print(f"Total paragraphs: {len(paras)}")
text_paras = [p for p in paras if p.text.strip()]
print(f"Non-empty paragraphs: {len(text_paras)}")

for i, p in enumerate(text_paras):
    txt = p.text[:100]
    print(f"  [{i}] style={p.style.name} | {txt}")

# Tables
tables = doc.tables
print(f"\nTables: {len(tables)}")
for ti, table in enumerate(tables):
    rows = len(table.rows)
    cols = len(table.columns) if table.rows else 0
    print(f"  Table {ti}: {rows}x{cols}")

# Sections
sections = doc.sections
print(f"\nSections: {len(sections)}")

# Check headers/footers
for si, section in enumerate(sections):
    header = section.header
    footer = section.footer
    h_paras = [p for p in header.paragraphs if p.text.strip()]
    f_paras = [p for p in footer.paragraphs if p.text.strip()]
    if h_paras or f_paras:
        print(f"  Section {si}: header={len(h_paras)} paras, footer={len(f_paras)} paras")
