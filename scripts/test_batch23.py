"""批2/批3 修复的离线验证（假翻译器，不调 API）。

验证点：
1. DOCX 嵌套表格 / 合并单元格去重 / 超链接 / 块级内容控件(sdt) / 文本框(txbxContent)
2. 对照模式换行转 w:br
3. TXT GBK 解码
4. PPTX 常规路径回归（SmartArt part 收集函数可运行）

用法（在 backend 镜像容器内）：python3 /scripts/test_batch23.py
"""
import io
import sys

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.models.task import OutputMode
from app.services.document_engine import (
    TranslationContext,
    _translate_docx_inplace,
    translate_txt,
    translate_pptx,
)

PASS = []
FAIL = []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(("✅" if cond else "❌"), name, detail if not cond else "")


class DummyTranslator:
    def __init__(self):
        self.calls = []

    def translate(self, text, target_lang, source_lang="auto", glossary=None, tm_reference=None):
        self.calls.append(text)
        return f"[T]{text}"


def build_docx() -> Document:
    doc = Document()
    doc.add_paragraph("Body paragraph one.")

    # 表格：嵌套表格 + 合并单元格
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Cell A"
    nested = t.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    merged = t.cell(1, 0).merge(t.cell(1, 1))
    merged.text = "Merged cell"

    # 超链接段落：普通 run + hyperlink run + 普通 run
    p = doc.add_paragraph()
    p.add_run("See ")
    p._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w")} w:anchor="top">'
        f'<w:r><w:t>the website</w:t></w:r></w:hyperlink>'
    ))
    p.add_run(" for details.")

    # 块级内容控件（sdt）
    doc.element.body.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent>'
        f'<w:p><w:r><w:t>SDT paragraph text</w:t></w:r></w:p>'
        f'</w:sdtContent></w:sdt>'
    ))

    # 文本框（VML w:pict / txbxContent），宿主段落自带文字
    doc.element.body.append(parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:t>Host paragraph.</w:t></w:r>'
        f'<w:r><w:pict xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<v:shape><v:textbox><w:txbxContent>'
        f'<w:p><w:r><w:t>Textbox text</w:t></w:r></w:p>'
        f'</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
    ))
    return doc


def test_docx_plain():
    doc = build_docx()
    tr = DummyTranslator()
    ctx = TranslationContext(
        target_lang="zh", source_lang="en",
        output_mode=OutputMode.PLAIN, translate_images="no",
    )
    _translate_docx_inplace(doc, tr, ctx)
    buf = io.BytesIO()
    doc.save(buf)
    doc2 = Document(io.BytesIO(buf.getvalue()))

    body_texts = [p.text for p in doc2.paragraphs]
    check("正文段落翻译", "[T]Body paragraph one." in body_texts)

    table = doc2.tables[0]
    check("普通单元格", table.cell(0, 0).text.strip() == "[T]Cell A")
    check("嵌套表格单元格", table.cell(0, 1).tables[0].cell(0, 0).text.strip() == "[T]Nested cell")
    check("合并单元格译文", table.cell(1, 0).text.strip() == "[T]Merged cell")
    check("合并单元格只翻一次", tr.calls.count("Merged cell") == 1,
          f"实际 {tr.calls.count('Merged cell')} 次")

    hyper_para = [p for p in doc2.paragraphs if "See" in p.text]
    check("超链接段落存在", len(hyper_para) == 1)
    if hyper_para:
        full = hyper_para[0].text
        check("超链接文字纳入翻译", full == "[T]See the website for details.",
              f"实际: {full!r}")
        check("超链接原文不残留", "the website" not in full.replace("[T]See the website for details.", ""))

    # 注意：docx 是 ZIP，必须查解析后的 body XML，不能对压缩字节做字符串搜索
    raw = doc2.element.body.xml
    check("内容控件(sdt)翻译", "[T]SDT paragraph text" in raw)
    check("文本框内容翻译", "[T]Textbox text" in raw)
    check("文本框宿主段落翻译", "[T]Host paragraph." in raw)
    check("宿主段落不含文本框文字", "[T]Host paragraph.Textbox" not in raw
          and "[T]Host paragraph.[T]Textbox" not in raw)
    check("文本框文字只翻一次", tr.calls.count("Textbox text") == 1,
          f"实际 {tr.calls.count('Textbox text')} 次")


def test_docx_bilingual_br():
    doc = Document()
    doc.add_paragraph("Hello world.")
    tr = DummyTranslator()
    ctx = TranslationContext(
        target_lang="zh", source_lang="en",
        output_mode=OutputMode.BILINGUAL, translate_images="no",
    )
    _translate_docx_inplace(doc, tr, ctx)
    p_xml = doc.paragraphs[0]._p.xml
    check("对照模式换行转 w:br", "<w:br" in p_xml and "Hello world." in p_xml, p_xml[:300])


def test_txt_gbk():
    data = "合同编号：ABC-123\n\n本协议由双方签署。".encode("gbk")
    tr = DummyTranslator()
    ctx = TranslationContext(
        target_lang="en", source_lang="zh",
        output_mode=OutputMode.PLAIN, translate_images="no",
    )
    out = translate_txt(data, tr, ctx).decode("utf-8")
    check("GBK 文本正确解码", "�" not in out and "本协议由双方签署" in out, out[:120])


def test_pptx_regression():
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly Review"
    buf = io.BytesIO()
    prs.save(buf)
    tr = DummyTranslator()
    ctx = TranslationContext(
        target_lang="zh", source_lang="en",
        output_mode=OutputMode.PLAIN, translate_images="no",
    )
    out = translate_pptx(buf.getvalue(), tr, ctx)
    prs2 = Presentation(io.BytesIO(out))
    title = prs2.slides[0].shapes.title.text
    check("PPTX 常规路径回归", title == "[T]Quarterly Review", title)


if __name__ == "__main__":
    test_docx_plain()
    test_docx_bilingual_br()
    test_txt_gbk()
    test_pptx_regression()
    print(f"\n结果: {len(PASS)} 通过, {len(FAIL)} 失败")
    sys.exit(1 if FAIL else 0)
