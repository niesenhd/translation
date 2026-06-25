"""调试 DOCX 段落内容"""
import sys
sys.path.insert(0, '/app')

from io import BytesIO
from minio import Minio
from app.core.database import SessionLocal
from app.models.task import TranslationTask
from app.core.config import get_settings
from docx import Document

settings = get_settings()
db = SessionLocal()
task = db.get(TranslationTask, 'ae9cc53f-f275-46b3-88b9-5c58aac29288')
print(f'file: {task.original_filename}')

mc = Minio(settings.minio_endpoint, access_key=settings.minio_access_key,
           secret_key=settings.minio_secret_key, secure=False)
obj = mc.get_object(settings.minio_bucket, f'tasks/{task.id}/original')
data = obj.read()

doc = Document(BytesIO(data))
empty_count = 0
short_count = 0
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        empty_count += 1
    elif len(text) < 3:
        short_count += 1
        if short_count <= 10:
            print(f'  short para[{i}]: len={len(text)} repr={repr(text)}')

print(f'Empty: {empty_count}, Short(<3): {short_count}, Total: {len(doc.paragraphs)}')

# 检查表格里的文本
table_count = len(doc.tables)
total_cells = sum(len(t.rows) * len(t.columns) for t in doc.tables)
print(f'Tables: {table_count}, Total cells: {total_cells}')
db.close()
